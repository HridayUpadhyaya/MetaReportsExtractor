from __future__ import annotations
from copy import copy
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference, Series

TEMPLATE = Path("data/meta_india_template.xlsx")
MASTER_HEADERS = [
    "Month", "Period Start", "Period End", "Report Published", "Platform", "Policy Category",
    "Content Actioned (raw)", "Content Actioned (numeric)", "Proactive Rate (%)",
    "Proactive Actions by AI (est.)", "Total User Grievances Received (platform-month)"
]


def _copy_row_style(ws, src_row: int, dst_row: int, max_col: int):
    for c in range(1, max_col + 1):
        src, dst = ws.cell(src_row, c), ws.cell(dst_row, c)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)


def build_workbook(rows: list[dict], failed_reports: list[str] | None = None) -> bytes:
    wb = load_workbook(TEMPLATE)
    failed_reports = failed_reports or []
    rows = sorted(rows, key=lambda r: (r.get("period_end") or "", r["platform"], r["policy_category"]))

    md = wb["Master Data"]
    # Preserve a copy of the template's first data-row style before clearing old data.
    style_cells = [copy(md.cell(2, c)._style) for c in range(1, 12)]
    style_numfmts = [md.cell(2, c).number_format for c in range(1, 12)]
    md.delete_rows(2, max(1, md.max_row - 1))
    for i, row in enumerate(rows, start=2):
        for c in range(1, 12):
            md.cell(i, c)._style = copy(style_cells[c-1])
            md.cell(i, c).number_format = style_numfmts[c-1]
        values = [
            row.get("month"), row.get("period_start"), row.get("period_end"), row.get("report_published"),
            row.get("platform"), row.get("policy_category"), row.get("content_actioned_raw"),
            row.get("content_actioned_numeric"), row.get("proactive_rate"), None, row.get("total_user_grievances"),
        ]
        for c, v in enumerate(values, start=1): md.cell(i, c).value = v
        md.cell(i, 10).value = f"=H{i}*I{i}"
    note_row = len(rows) + 3
    md.cell(note_row, 1).value = "Note: 'Total User Grievances Received' is the platform-month total (not broken down by category) and is repeated on every category row for that platform-month, so it can be filtered/pivoted alongside category-level data. 'Proactive Actions by AI (est.)' = Content Actioned x Proactive Rate."
    md.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=11)
    md.freeze_panes = "A2"
    md.auto_filter.ref = f"A1:K{max(2, len(rows)+1)}"
    md.column_dimensions["F"].width = 60
    md.column_dimensions["K"].width = 28

    # Monthly Summary
    ms = wb["Monthly Summary"]
    if ms.max_row > 1: ms.delete_rows(2, ms.max_row - 1)
    combos = sorted({(r["month"], r.get("period_end"), r["platform"]) for r in rows})
    master_end = len(rows) + 1
    for i, (month, period_end, platform) in enumerate(combos, start=2):
        _copy_row_style(ms, 2 if ms.max_row >= 2 else 1, i, 7)
        ms.cell(i,1).value, ms.cell(i,2).value, ms.cell(i,3).value = month, period_end, platform
        ms.cell(i,4).value = f'=SUMIFS(\'Master Data\'!$H$2:$H${master_end},\'Master Data\'!$A$2:$A${master_end},A{i},\'Master Data\'!$E$2:$E${master_end},C{i})'
        ms.cell(i,5).value = f'=SUMIFS(\'Master Data\'!$J$2:$J${master_end},\'Master Data\'!$A$2:$A${master_end},A{i},\'Master Data\'!$E$2:$E${master_end},C{i})'
        ms.cell(i,6).value = f'=IFERROR(E{i}/D{i},"")'
        ms.cell(i,7).value = f'=IFERROR(AVERAGEIFS(\'Master Data\'!$K$2:$K${master_end},\'Master Data\'!$A$2:$A${master_end},A{i},\'Master Data\'!$E$2:$E${master_end},C{i}),"")'
    ms.freeze_panes = "A2"

    # Trend Charts data + two charts.
    tc = wb["Trend Charts"]
    tc._charts = []
    if tc.max_row > 1: tc.delete_rows(2, tc.max_row - 1)
    months = sorted({(r["month"], r.get("period_end")) for r in rows})
    for i, (month, period_end) in enumerate(months, start=2):
        tc.cell(i,1).value, tc.cell(i,2).value = month, period_end
        tc.cell(i,3).value = f'=SUMIFS(\'Master Data\'!$H$2:$H${master_end},\'Master Data\'!$A$2:$A${master_end},A{i},\'Master Data\'!$E$2:$E${master_end},"Facebook")'
        tc.cell(i,4).value = f'=IFERROR(AVERAGEIFS(\'Master Data\'!$K$2:$K${master_end},\'Master Data\'!$A$2:$A${master_end},A{i},\'Master Data\'!$E$2:$E${master_end},"Facebook"),"")'
        tc.cell(i,5).value = f'=SUMIFS(\'Master Data\'!$H$2:$H${master_end},\'Master Data\'!$A$2:$A${master_end},A{i},\'Master Data\'!$E$2:$E${master_end},"Instagram")'
        tc.cell(i,6).value = f'=IFERROR(AVERAGEIFS(\'Master Data\'!$K$2:$K${master_end},\'Master Data\'!$A$2:$A${master_end},A{i},\'Master Data\'!$E$2:$E${master_end},"Instagram"),"")'
    if months:
        end = len(months) + 1
        cats = Reference(tc, min_col=1, min_row=2, max_row=end)
        c1 = LineChart(); c1.title = "Content Actioned Trend"; c1.y_axis.title = "Content actioned"; c1.x_axis.title = "Month"
        for col in (3, 5):
            series = Series(Reference(tc, min_col=col, min_row=2, max_row=end), title_from_data=False, title=tc.cell(1,col).value)
            c1.series.append(series)
        c1.set_categories(cats); c1.height=8; c1.width=16
        tc.add_chart(c1, "H1")
        c2 = LineChart(); c2.title = "User Grievances Received Trend"; c2.y_axis.title = "Grievances"; c2.x_axis.title = "Month"
        for col in (4, 6):
            series = Series(Reference(tc, min_col=col, min_row=2, max_row=end), title_from_data=False, title=tc.cell(1,col).value)
            c2.series.append(series)
        c2.set_categories(cats); c2.height=8; c2.width=16
        tc.add_chart(c2, "H26")

    notes = wb["Data Notes"]
    periods = [r.get("period_end") for r in rows if r.get("period_end")]
    platforms = sorted({r["platform"] for r in rows})
    review_count = sum(1 for r in rows if r.get("validation_status") != "ok")
    if periods:
        notes["A11"] = f"{len(set(r['month'] for r in rows))} monthly data points, {min(periods)} – {max(periods)}. Platforms present: {', '.join(platforms)}."
    notes["A7"] = f"{review_count} extracted row(s) are flagged for review by validation rules."
    notes["A8"] = f"{len(failed_reports)} report(s) could not be parsed/downloaded in the latest run."

    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    except Exception:
        pass
    bio = BytesIO(); wb.save(bio); return bio.getvalue()
